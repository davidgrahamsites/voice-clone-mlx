"""Parse source files into normalized text."""

import zipfile
from pathlib import Path
from typing import List

from voiceclonegpt.ingestion import web

# Resource limits for .docx ingestion. A .docx is a ZIP archive, so an untrusted
# file can be a zip bomb; these bound the work before anything is decompressed.
MAX_DOCX_FILE_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 2000
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200
MAX_DOCX_TEXT_CHARS = 5_000_000

# Resource limits for .pdf ingestion.
MAX_PDF_FILE_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 2000
MAX_PDF_TEXT_CHARS = 5_000_000


def parse_source_file(file_path: Path) -> str:
    """Parse a source file and extract text content.

    Args:
        file_path: Path to source file

    Returns:
        Extracted text content

    Raises:
        FileNotFoundError: If file doesn't exist
        ValueError: If file format is not supported
    """
    if web.is_url(file_path):
        return web.fetch_url(str(file_path))

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Source file not found: {file_path}")

    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        return file_path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        return _parse_docx(file_path)
    elif suffix == ".pdf":
        return _parse_pdf(file_path)
    else:
        raise ValueError(
            f"Unsupported file format: {suffix}. "
            f"Supported: .txt, .docx, .pdf"
        )


def _preflight_docx(file_path: Path) -> None:
    """Check a .docx against resource limits before decompressing it.

    Raises:
        ValueError: If the file exceeds a limit or is not a readable package
    """
    size = file_path.stat().st_size
    if size > MAX_DOCX_FILE_BYTES:
        raise ValueError(
            f"DOCX file too large: {size} bytes "
            f"(limit {MAX_DOCX_FILE_BYTES})"
        )

    try:
        with zipfile.ZipFile(file_path) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError(
            f"File is not a readable DOCX package: {file_path}"
        ) from exc

    if len(entries) > MAX_DOCX_ENTRIES:
        raise ValueError(
            f"DOCX package has too many entries: {len(entries)} "
            f"(limit {MAX_DOCX_ENTRIES})"
        )

    uncompressed = sum(entry.file_size for entry in entries)
    if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise ValueError(
            f"DOCX uncompressed contents too large: {uncompressed} bytes "
            f"(limit {MAX_DOCX_UNCOMPRESSED_BYTES})"
        )

    compressed = sum(entry.compress_size for entry in entries)
    ratio = uncompressed / max(compressed, 1)
    if ratio > MAX_DOCX_COMPRESSION_RATIO:
        raise ValueError(
            f"DOCX compression ratio too high: {ratio:.0f}x "
            f"(limit {MAX_DOCX_COMPRESSION_RATIO}x)"
        )


def _parse_docx(file_path: Path) -> str:
    """Extract paragraph and table text from a .docx file.

    Raises:
        ImportError: If python-docx is not installed
        ValueError: If the file exceeds a resource limit, is malformed, or
            yields more text than the extraction budget allows
    """
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "Reading .docx files requires the python-docx package. "
            "Install it with: pip install python-docx"
        ) from exc

    _preflight_docx(file_path)

    try:
        document = docx.Document(str(file_path))
    except Exception as exc:
        raise ValueError(
            f"File is not a readable DOCX package: {file_path}"
        ) from exc

    lines = [p.text.strip() for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text.strip())

    text = "\n".join(line for line in lines if line)

    if len(text) > MAX_DOCX_TEXT_CHARS:
        raise ValueError(
            f"DOCX extracted text too large: {len(text)} characters "
            f"(limit {MAX_DOCX_TEXT_CHARS})"
        )

    return text


def _default_pdf_reader(file_path: Path):
    """Open a PDF with pypdf, the declared optional dependency.

    Args:
        file_path: Path to the .pdf file

    Returns:
        An object exposing `.pages`, each with `.extract_text()`

    Raises:
        ImportError: If pypdf is not installed
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "Reading .pdf files requires the pypdf package. "
            "Install it with: pip install pypdf"
        ) from exc

    return PdfReader(str(file_path))


def _parse_pdf(file_path: Path, reader_factory=None) -> str:
    """Extract text from a text-based PDF.

    Scanned (image-only) PDFs are rejected rather than silently returning
    nothing; OCR is out of scope.

    Args:
        file_path: Path to the .pdf file
        reader_factory: Callable returning a reader for the path. Defaults to
            `_default_pdf_reader` (pypdf); injectable for tests.

    Returns:
        Extracted text, one non-empty page per line group

    Raises:
        ImportError: If pypdf is needed but not installed
        ValueError: If the file exceeds a limit, is malformed, has no pages,
            or contains no extractable text
    """
    size = file_path.stat().st_size
    if size > MAX_PDF_FILE_BYTES:
        raise ValueError(
            f"PDF file too large: {size} bytes (limit {MAX_PDF_FILE_BYTES})"
        )

    if reader_factory is None:
        reader_factory = _default_pdf_reader

    try:
        reader = reader_factory(file_path)
        pages = list(reader.pages)
    except ImportError:
        raise
    except Exception as exc:
        # Includes ValueError from the PDF library: normalize it, since a
        # library ValueError means "unreadable file", not one of our limits.
        raise ValueError(
            f"File is not a readable PDF: {file_path}"
        ) from exc

    if not pages:
        raise ValueError(f"PDF has no pages: {file_path}")

    if len(pages) > MAX_PDF_PAGES:
        raise ValueError(
            f"PDF has too many pages: {len(pages)} (limit {MAX_PDF_PAGES})"
        )

    # Accumulate page by page so a document that blows the text budget is
    # abandoned at the offending page, not after extracting all of it.
    chunks = []
    total_chars = 0

    for page in pages:
        try:
            page_text = page.extract_text()
        except Exception as exc:
            raise ValueError(
                f"File is not a readable PDF: {file_path}"
            ) from exc

        if not page_text or not page_text.strip():
            continue

        chunk = page_text.strip()
        chunks.append(chunk)
        total_chars += len(chunk) + (1 if len(chunks) > 1 else 0)

        # Raised outside the extraction try so it is never normalized away.
        if total_chars > MAX_PDF_TEXT_CHARS:
            raise ValueError(
                f"PDF extracted text too large: exceeded "
                f"{MAX_PDF_TEXT_CHARS} characters"
            )

    text = "\n".join(chunks)

    if not text:
        raise ValueError(
            f"PDF contains no extractable text (it is likely scanned): "
            f"{file_path}. OCR is not supported; supply a text-based PDF, "
            f"or export the document as .txt or .docx."
        )

    return text


def split_into_sentences(text: str) -> List[str]:
    """Split text into sentences.

    Args:
        text: Input text

    Returns:
        List of sentences
    """
    import re

    # Simple sentence splitter - split on . ! ? followed by space
    sentences = re.split(r'(?<=[.!?])\s+', text.strip())
    return [s.strip() for s in sentences if s.strip()]
